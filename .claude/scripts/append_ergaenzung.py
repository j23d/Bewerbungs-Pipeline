# -*- coding: utf-8 -*-
"""Stufe 3: Schreibt eine im Luecken-Dialog gesicherte Antwort dauerhaft nach
profil/ergaenzungen.docx zurueck, damit dieselbe Luecke kuenftig nicht erneut
abgefragt wird.

Aufruf:
  python .claude/scripts/append_ergaenzung.py \
    --datum 2026-07-01 --lauf runs/2026-07-01-1200-acme \
    --frage "..." --antwort "..."
"""
import argparse
import os
from docx import Document

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
ERG = os.path.join(ROOT, "profil", "ergaenzungen.docx")


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--datum", required=True)
    p.add_argument("--lauf", required=True)
    p.add_argument("--frage", required=True)
    p.add_argument("--antwort", required=True)
    args = p.parse_args()

    doc = Document(ERG) if os.path.exists(ERG) else Document()
    doc.add_paragraph("")
    h = doc.add_paragraph()
    r = h.add_run(f"[{args.datum} | Lauf: {args.lauf}]")
    r.bold = True
    doc.add_paragraph(f"Frage: {args.frage}")
    doc.add_paragraph(f"Antwort: {args.antwort}")
    doc.save(ERG)
    print(f"Ergaenzung gespeichert in {ERG}")


if __name__ == "__main__":
    main()

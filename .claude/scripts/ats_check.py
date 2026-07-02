# -*- coding: utf-8 -*-
"""Stufe 5b: ATS-Simulator.

Liest die FERTIG gerenderten DOCX-Dateien eines Laufs zurueck – so, wie es ein
Applicant Tracking System (ATS) beim Parsen taete – und liefert einen
UEBERPRUEFBAREN (nicht geschaetzten) ATS-Befund:

  * gemessene Keyword-Abdeckung der Must-haves aus 01_analyse.md,
  * Parsing-Warnungen (kritischer Text in Tabellen/Textboxen/Kopf-Fusszeilen,
    fehlende Standard-Ueberschriften, Kontaktdaten nicht im Fliesstext),
  * eine .txt-Extraktion als Fallback-Ansicht dessen, was das ATS "sieht".

WICHTIG (Kontrollfluss): Dies ist ein WARNUNG-UND-FREIGABE-Schritt, KEIN hartes
Gate. Liegt die Abdeckung unter der Schwelle, wird laut gewarnt – blockiert wird
nichts. Die Entscheidung trifft der Mensch bei der Freigabe.

Aufruf:
  python .claude/scripts/ats_check.py \
    --analyse runs/<lauf>/01_analyse.md \
    --run runs/<lauf> \
    [--schwelle 75]

Ergebnis:
  * schreibt <run>/05b_ats.json
  * schreibt <run>/ats_extrakt.txt (das, was das ATS liest)
  * Exit 0 = ueber Schwelle, Exit 3 = unter Schwelle (Warnung, NICHT blockierend)
"""
import argparse
import json
import os
import re
import unicodedata

from docx import Document

# Standard-Ueberschriften, die ein ATS zuverlaessig erkennt.
STANDARD_HEADINGS = [
    "berufserfahrung", "erfahrung", "ausbildung", "kenntnisse", "skills",
    "qualifikationen", "kurzprofil", "profil",
]


# --------------------------------------------------------------------------- #
# DOCX-Textextraktion – bewusst so, wie ein ATS parst.
# --------------------------------------------------------------------------- #
def iter_block_text(doc):
    """Fliesstext aus normalen Absaetzen (die 'sichere' Zone)."""
    return [p.text for p in doc.paragraphs if p.text.strip()]


def iter_table_text(doc):
    """Text in Tabellen – von vielen ATS falsch oder gar nicht gelesen."""
    out = []
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text.strip():
                    out.append(cell.text.strip())
    return out


def iter_header_footer_text(doc):
    """Text in Kopf-/Fusszeilen – kritische Infos hier gehen oft verloren."""
    out = []
    for sec in doc.sections:
        for container in (sec.header, sec.footer,
                          sec.first_page_header, sec.first_page_footer):
            for p in container.paragraphs:
                if p.text.strip():
                    out.append(p.text.strip())
    return out


def iter_textbox_text(doc):
    """Text in Textboxen/Shapes (w:txbxContent) – fuer ATS meist unsichtbar."""
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    out = []
    body = doc.element.body
    for tb in body.iter("{%s}txbxContent" % ns["w"]):
        txt = "".join(node.text or "" for node in tb.iter("{%s}t" % ns["w"]))
        if txt.strip():
            out.append(txt.strip())
    return out


def extract(docx_path):
    """Gibt (safe_text, risky_text, warnings) zurueck."""
    doc = Document(docx_path)
    body = iter_block_text(doc)
    tables = iter_table_text(doc)
    headers = iter_header_footer_text(doc)
    boxes = iter_textbox_text(doc)

    warnings = []
    if tables:
        warnings.append(
            f"{os.path.basename(docx_path)}: Text in Tabellen gefunden "
            f"({len(tables)} Zellen). Viele ATS lesen Tabellen fehlerhaft – "
            f"kritische Inhalte besser als Fliesstext.")
    if headers:
        warnings.append(
            f"{os.path.basename(docx_path)}: Text in Kopf-/Fusszeile gefunden. "
            f"Kontaktdaten/kritische Infos gehoeren in den Fliesstext, nicht dorthin.")
    if boxes:
        warnings.append(
            f"{os.path.basename(docx_path)}: Text in Textboxen/Shapes gefunden. "
            f"Diese sind fuer ATS meist unsichtbar.")

    safe = "\n".join(body)
    risky = "\n".join(tables + headers + boxes)
    return safe, risky, warnings


# --------------------------------------------------------------------------- #
# Keyword-Extraktion aus 01_analyse.md
# --------------------------------------------------------------------------- #
def parse_must_haves_and_keywords(analyse_md):
    """Zieht Must-haves und ATS-Keywords aus den entsprechenden Abschnitten."""
    def section(title):
        m = re.search(rf"^#+\s*{re.escape(title)}.*?$(.*?)(?=^#+\s|\Z)",
                      analyse_md, re.MULTILINE | re.DOTALL | re.IGNORECASE)
        return m.group(1) if m else ""

    def bullets(block):
        items = []
        for line in block.splitlines():
            line = line.strip()
            line = re.sub(r"^\d+\.\s*", "", line)   # "1. ..."
            line = re.sub(r"^[-*]\s*", "", line)     # "- ..."
            if line and not line.startswith("#"):
                items.append(line)
        return items

    must = bullets(section("Must-haves"))
    kw_block = section("Keywords")
    # Keywords sind oft komma-separiert je Zeile -> aufsplitten.
    keywords = []
    for item in bullets(kw_block):
        keywords.extend([k.strip() for k in re.split(r"[,/]", item) if k.strip()])
    return must, keywords


# --------------------------------------------------------------------------- #
# Matching
# --------------------------------------------------------------------------- #
def norm(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return s.lower()


def extract_terms(phrase):
    """Aus einer Anforderung/einem Keyword die eigentlichen Suchbegriffe ziehen.
    Klammern (z.B. 'ERP (DIAMANT, SAP)') werden als einzelne Terme mitgenommen."""
    terms = set()
    # Begriffe in Klammern separat.
    for grp in re.findall(r"\(([^)]*)\)", phrase):
        for t in re.split(r"[,/]", grp):
            t = t.strip()
            if len(t) >= 3:
                terms.add(t)
    base = re.sub(r"\([^)]*\)", " ", phrase)
    # Signifikante Woerter (>=4 Zeichen) als Fallback-Terme.
    for w in re.findall(r"[A-Za-zÄÖÜäöüß\-]{4,}", base):
        terms.add(w)
    return terms


STOP = set(norm(w) for w in [
    "oder", "und", "sowie", "eine", "einen", "einer", "abgeschlossenes",
    "abgeschlossene", "erfahrung", "kenntnisse", "interesse", "sehr", "gute",
    "mit", "der", "die", "das", "fuer", "als", "bzw", "z.b", "zum", "zur",
])


def covered(term, haystack):
    return norm(term) in haystack


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--analyse", required=True)
    ap.add_argument("--run", required=True)
    ap.add_argument("--schwelle", type=int, default=75)
    args = ap.parse_args()

    with open(args.analyse, encoding="utf-8") as f:
        analyse_md = f.read()
    must_haves, keywords = parse_must_haves_and_keywords(analyse_md)

    # Alle DOCX des Laufs einlesen (Anschreiben + Lebenslauf).
    docx_files = [os.path.join(args.run, fn) for fn in os.listdir(args.run)
                  if fn.lower().endswith(".docx")]
    docx_files.sort()

    all_safe, all_risky, warnings = [], [], []
    for d in docx_files:
        safe, risky, warns = extract(d)
        all_safe.append(safe)
        all_risky.append(risky)
        warnings.extend(warns)

    safe_text = norm("\n".join(all_safe))
    risky_text = norm("\n".join(all_risky))
    full_text = safe_text + "\n" + risky_text

    # Extrakt fuer den Menschen sichern (was das ATS "liest").
    with open(os.path.join(args.run, "ats_extrakt.txt"), "w", encoding="utf-8") as f:
        f.write("=== ATS-SICHT: Fliesstext (zuverlaessig geparst) ===\n")
        f.write("\n".join(all_safe))
        f.write("\n\n=== RISIKO-ZONE: Tabellen/Kopfzeilen/Textboxen ===\n")
        f.write("\n".join(all_risky) or "(keine)")

    # --- Must-have-Abdeckung messen ---
    mh_results = []
    for mh in must_haves:
        terms = {t for t in extract_terms(mh) if norm(t) not in STOP}
        hits = [t for t in terms if covered(t, full_text)]
        # Anforderung gilt als abgedeckt, wenn mind. ein signifikanter Term auftaucht.
        mh_results.append({
            "anforderung": mh,
            "abgedeckt": bool(hits),
            "treffer": sorted(set(hits)),
            "nur_in_risikozone": bool(hits) and not any(covered(t, safe_text) for t in hits),
        })
    mh_covered = sum(1 for r in mh_results if r["abgedeckt"])
    mh_score = round(100 * mh_covered / len(mh_results)) if mh_results else 0

    # --- Keyword-Abdeckung messen ---
    kw_results = []
    for kw in keywords:
        kw_results.append({"keyword": kw, "gefunden": covered(kw, full_text),
                           "nur_in_risikozone": covered(kw, risky_text) and not covered(kw, safe_text)})
    kw_covered = sum(1 for r in kw_results if r["gefunden"])
    kw_score = round(100 * kw_covered / len(kw_results)) if kw_results else 0

    # Gesamt-ATS-Score: Must-haves doppelt gewichtet.
    gesamt = round((2 * mh_score + kw_score) / 3)

    # Zusaetzliche strukturelle Warnungen.
    if not any(h in safe_text for h in STANDARD_HEADINGS):
        warnings.append("Keine der erwarteten Standard-Ueberschriften "
                        "(Berufserfahrung/Ausbildung/Kenntnisse ...) im Fliesstext "
                        "gefunden – ATS erkennt die Abschnitte evtl. nicht.")
    if not re.search(r"@", "\n".join(all_safe)):
        warnings.append("Keine E-Mail-Adresse im Fliesstext gefunden – "
                        "Kontaktdaten evtl. nicht ATS-lesbar (z.B. in Kopfzeile?).")
    for r in mh_results:
        if r["nur_in_risikozone"]:
            warnings.append(f"Must-have '{r['anforderung'][:60]}' nur in Risiko-Zone "
                            f"(Tabelle/Kopfzeile/Textbox) gefunden – fuer ATS evtl. verloren.")

    bestanden = gesamt >= args.schwelle
    result = {
        "ats_score_gemessen": gesamt,
        "schwelle": args.schwelle,
        "ueber_schwelle": bestanden,
        "blockiert": False,  # Warnung-und-Freigabe: NIE blockierend
        "must_have_score": mh_score,
        "keyword_score": kw_score,
        "must_haves": mh_results,
        "keywords": kw_results,
        "geprueft_dateien": [os.path.basename(d) for d in docx_files],
        "parsing_warnungen": warnings,
    }
    with open(os.path.join(args.run, "05b_ats.json"), "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    # Konsolen-Zusammenfassung.
    print(f"ATS-Score (gemessen): {gesamt}/100  (Schwelle {args.schwelle})")
    print(f"  Must-haves: {mh_covered}/{len(mh_results)} abgedeckt ({mh_score}%)")
    print(f"  Keywords:   {kw_covered}/{len(kw_results)} gefunden ({kw_score}%)")
    if warnings:
        print("Parsing-Warnungen:")
        for w in warnings:
            print("  - " + w)
    if not bestanden:
        print(f"\nWARNUNG: ATS-Score unter Schwelle. Dies BLOCKIERT NICHT – "
              f"du entscheidest bei der Freigabe (Warnung-und-Freigabe).")
    print(f"\nDetails: {os.path.join(args.run, '05b_ats.json')}")
    print(f"ATS-Sicht: {os.path.join(args.run, 'ats_extrakt.txt')}")

    # Exit 3 = unter Schwelle (Signal), aber ausdruecklich NICHT blockierend.
    raise SystemExit(0 if bestanden else 3)


if __name__ == "__main__":
    main()
